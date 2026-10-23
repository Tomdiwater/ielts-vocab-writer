# -*- coding: utf-8 -*-
"""
ielts-vocab-writer 公共模块
所有格式参数提取自模板原生词条，保证排版与模板一致。
"""
import os
import sys
import json
import shutil
from pathlib import Path

from docx import Document
from docx.shared import RGBColor, Pt, Emu

# ============ 格式常量（与模板一致，勿随意修改） ============
COLOR_HEADWORD = RGBColor(0x00, 0x00, 0x00)   # 词头：黑色
COLOR_SUPP = RGBColor(0x50, 0x50, 0x50)       # 补充行：深灰
SUPP_INDENT = Emu(274320)                      # 补充行左缩进
SUPP_SIZE = Pt(9)                              # 补充行字号 9pt
MAIN_SPACE_BEFORE = Emu(63500)                 # 主行段前 5pt
MAIN_SPACE_AFTER = Emu(12700)                  # 主行段后 1pt
SUPP_SPACE_AFTER = Emu(63500)                  # 补充行段后 5pt
HEADING_SIZE = Pt(13)                          # 章节/子标题字号
HEADING_SPACE_BEFORE = Emu(177800)             # 章节标题段前 14pt
HEADING_SPACE_AFTER = Emu(76200)               # 章节标题段后 6pt

DEFAULT_DOC_NAME = '雅思阅读词汇.docx'
CN_NUM = '一二三四五六七八九十'


# ============ 文档定位与初始化 ============
def get_desktop_path():
    """跨平台获取桌面路径（Windows 优先读注册表，兼容中文系统重定向）"""
    if sys.platform.startswith('win'):
        try:
            import winreg
            key = winreg.OpenKey(
                winreg.HKEY_CURRENT_USER,
                r'Software\Microsoft\Windows\CurrentVersion\Explorer\Shell Folders')
            desktop, _ = winreg.QueryValueEx(key, 'Desktop')
            if desktop and os.path.isdir(desktop):
                return desktop
        except Exception:
            pass
    return str(Path.home() / 'Desktop')


# ============ 用户配置（默认路径可永久自定义） ============
CONFIG_DIR = Path.home() / '.ielts-vocab-writer'
CONFIG_FILE = CONFIG_DIR / 'config.json'


def load_config():
    """读取用户配置 dict；不存在或损坏时返回 {}"""
    try:
        if CONFIG_FILE.exists():
            with open(CONFIG_FILE, 'r', encoding='utf-8') as f:
                return json.load(f)
    except Exception:
        pass
    return {}


def save_config(cfg):
    CONFIG_DIR.mkdir(parents=True, exist_ok=True)
    with open(CONFIG_FILE, 'w', encoding='utf-8') as f:
        json.dump(cfg, f, ensure_ascii=False, indent=2)


def get_default_doc_path():
    """解析默认文档路径：用户配置优先，否则桌面《雅思阅读词汇.docx》"""
    p = load_config().get('document_path')
    if p:
        return p
    return os.path.join(get_desktop_path(), DEFAULT_DOC_NAME)


def set_default_path(path):
    """永久保存默认文档路径。path 可以是目录（自动补文件名）或完整 .docx 路径。
    返回规范化后的完整路径。"""
    path = os.path.abspath(os.path.expanduser(path))
    base = os.path.basename(path)
    # 已存在的目录、或没有文件扩展名的路径，一律视为目录
    if os.path.isdir(path) or ('.' not in base and not base.endswith('.docx')):
        path = os.path.join(path, DEFAULT_DOC_NAME)
    save_config({'document_path': path})
    return path


def reset_default_path():
    """清除自定义路径，恢复桌面默认。返回恢复后的完整路径。"""
    try:
        if CONFIG_FILE.exists():
            CONFIG_FILE.unlink()
    except Exception:
        pass
    return os.path.join(get_desktop_path(), DEFAULT_DOC_NAME)


def is_first_run():
    """首次使用：没有自定义配置，且桌面默认文档尚不存在"""
    cfg = load_config()
    if cfg.get('document_path'):
        return False
    return not os.path.exists(os.path.join(get_desktop_path(), DEFAULT_DOC_NAME))


def template_path():
    """返回 skill 内置模板路径"""
    return Path(__file__).resolve().parent.parent / 'assets' / 'template.docx'


def resolve_latest(doc_path):
    """若存在 _v2/_v3 等另存版本，返回版本号最大的那个（文件被占用时会另存）"""
    base, ext = os.path.splitext(doc_path)
    latest = doc_path
    i = 2
    while os.path.exists(f'{base}_v{i}{ext}'):
        latest = f'{base}_v{i}{ext}'
        i += 1
    return latest


def ensure_document(doc_path=None):
    """返回应写入的文档路径：--document 参数 > 用户配置 > 桌面默认；不存在则从模板初始化"""
    if doc_path is None:
        doc_path = get_default_doc_path()
    actual = resolve_latest(doc_path) if os.path.exists(doc_path) else doc_path
    if not os.path.exists(actual):
        tpl = template_path()
        if not tpl.exists():
            raise FileNotFoundError(f'模板文件不存在: {tpl}')
        os.makedirs(os.path.dirname(actual), exist_ok=True)
        shutil.copy2(tpl, actual)
        print(f'[init] 已从模板创建新文档: {actual}')
    return actual


# ============ 章节结构识别（动态读取，不写死分类） ============
def is_section_heading(p):
    t = p.text.strip()
    return len(t) >= 2 and t[0] in CN_NUM and t[1] == '、'


def is_subheading(p):
    t = p.text.strip()
    return t.startswith('【') and '】' in t


def find_section(doc, keyword):
    """按关键词模糊匹配章节标题，返回 (索引, 段落)；找不到返回 (None, None)"""
    for i, p in enumerate(doc.paragraphs):
        if is_section_heading(p) and keyword in p.text:
            return i, p
    return None, None


def next_section_paragraph(doc, start_idx):
    """返回 start_idx 之后的下一个章节标题段落（作为插入锚点）；无则 None（即文档末尾）"""
    for p in doc.paragraphs[start_idx + 1:]:
        if is_section_heading(p):
            return p
    return None


def find_subheading(doc, label):
    """查找【xx】子标题段落，返回 (索引, 段落)；找不到返回 (None, None)。
    label 可含或不含【】。"""
    target = label if label.startswith('【') else f'【{label}】'
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip() == target:
            return i, p
    return None, None


# ============ 词条写入（核心格式） ============
def _style_main(p_main):
    pf = p_main.paragraph_format
    pf.space_before = MAIN_SPACE_BEFORE
    pf.space_after = MAIN_SPACE_AFTER


def _style_supp(p_supp):
    pf = p_supp.paragraph_format
    pf.left_indent = SUPP_INDENT
    pf.space_after = SUPP_SPACE_AFTER


def _fill_runs(p_main, p_supp, head, head_rest, supplements):
    """填充主行和补充行内容"""
    r1 = p_main.add_run(head)
    r1.bold = True
    r1.font.color.rgb = COLOR_HEADWORD
    if head_rest:
        p_main.add_run(head_rest)
    _style_main(p_main)

    if supplements:
        for i, line in enumerate(supplements):
            r = p_supp.add_run(f'▸ {line}')
            r.font.size = SUPP_SIZE
            r.font.color.rgb = COLOR_SUPP
            if i < len(supplements) - 1:
                br = p_supp.add_run('\n')
                br.font.size = SUPP_SIZE
        _style_supp(p_supp)


def add_entry_at_anchor(anchor, head, head_rest, supplements):
    """在锚点段落前插入一个词条（主行 + 补充行）"""
    p_main = anchor.insert_paragraph_before('')
    p_supp = anchor.insert_paragraph_before('')
    _fill_runs(p_main, p_supp, head, head_rest, supplements)


def add_entry_at_end(doc, head, head_rest, supplements):
    """在文档末尾追加一个词条"""
    p_main = doc.add_paragraph('')
    p_supp = doc.add_paragraph('')
    _fill_runs(p_main, p_supp, head, head_rest, supplements)


# ============ 三类写入入口 ============
def add_word(doc, data):
    """写入单词到指定章节。data 字段：section/word/phonetic/pos_meaning/supplements"""
    section_kw = data.get('section', '通用')
    idx, heading = find_section(doc, section_kw)
    if heading is None:
        # 找不到章节时落到通用词汇
        idx, heading = find_section(doc, '通用')
    if heading is None:
        raise ValueError(f'找不到章节「{section_kw}」，且文档中无通用词汇章节')
    anchor = next_section_paragraph(doc, idx)
    head = data['word']
    rest_parts = []
    if data.get('phonetic'):
        rest_parts.append(data['phonetic'])
    if data.get('pos_meaning'):
        rest_parts.append(data['pos_meaning'])
    head_rest = '  ' + ' '.join(rest_parts) if rest_parts else ''
    if anchor is not None:
        add_entry_at_anchor(anchor, head, head_rest, data.get('supplements', []))
    else:
        add_entry_at_end(doc, head, head_rest, data.get('supplements', []))
    return f'单词「{head}」→ {heading.text.strip()}'


def add_phrase(doc, data):
    """写入短语到短语区。data 字段：phrase/meaning/supplements"""
    idx, heading = find_section(doc, '短语区')
    if heading is None:
        raise ValueError('文档中找不到「短语区」章节')
    anchor = next_section_paragraph(doc, idx)
    head = data['phrase']
    head_rest = f'  {data["meaning"]}' if data.get('meaning') else ''
    if anchor is not None:
        add_entry_at_anchor(anchor, head, head_rest, data.get('supplements', []))
    else:
        add_entry_at_end(doc, head, head_rest, data.get('supplements', []))
    return f'短语「{head}」→ 短语区'


def add_synonym(doc, data):
    """写入同义/反义词组。data 字段：relation(synonym|antonym)/terms(list)/meaning/supplements"""
    relation = data.get('relation', 'synonym')
    sub_label = '同义词组' if relation == 'synonym' else '反义词组'
    sub_idx, sub = find_subheading(doc, sub_label)
    if sub is None:
        raise ValueError(f'同义替换区找不到子标题「{sub_label}」')

    terms = data['terms']
    head = ' / '.join(terms)
    head_rest = f'  {data["meaning"]}' if data.get('meaning') else ''
    supplements = data.get('supplements', [])

    if relation == 'synonym':
        # 同义词组锚点 = 【反义词组】子标题；找不到则用下一个章节标题
        _, anchor = find_subheading(doc, '反义词组')
        if anchor is None:
            anchor = next_section_paragraph(doc, sub_idx)
        if anchor is not None:
            add_entry_at_anchor(anchor, head, head_rest, supplements)
        else:
            add_entry_at_end(doc, head, head_rest, supplements)
    else:
        # 反义词组：锚点为【反义词组】后面的章节标题（无则文末）；
        # "在锚点前插入"天然追加到反义词组区域末尾
        anchor = next_section_paragraph(doc, sub_idx)
        if anchor is not None:
            add_entry_at_anchor(anchor, head, head_rest, supplements)
        else:
            add_entry_at_end(doc, head, head_rest, supplements)
    return f'{sub_label}「{head}」'


# ============ 保存（文件占用自动另存） ============
def save_document(doc, path):
    try:
        doc.save(path)
        return path
    except PermissionError:
        base, ext = os.path.splitext(path)
        i = 2
        while os.path.exists(f'{base}_v{i}{ext}'):
            i += 1
        new_path = f'{base}_v{i}{ext}'
        doc.save(new_path)
        print(f'[warn] 原文件被占用（可能正在 Word 中打开），已另存为: {new_path}')
        return new_path
